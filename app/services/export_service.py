"""Excel and AI-friendly exports sourced only from SQLite queries."""

from __future__ import annotations

import csv
import json
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.database import Database


DETAIL_HEADERS = [
    "日期", "人员", "小组", "项目", "模块", "工作类别", "具体工作", "规则ID",
    "实际数量", "实际单位", "标准化数量", "定额单位", "基准分", "工作量", "备注",
]


class ExportService:
    def __init__(self, database: Database):
        self.database = database

    def export_day(self, work_date: str, target: str | Path) -> Path:
        try:
            date.fromisoformat(work_date)
        except ValueError as exc:
            raise ValueError("日期必须采用有效的 YYYY-MM-DD 格式") from exc
        rows = self.database.list_records(date_from=work_date, date_to=work_date)
        path = Path(target)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "工作量明细"
        self._write_detail_sheet(sheet, rows)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path

    def export_month(self, month: str, target: str | Path) -> Path:
        if not re.fullmatch(r"\d{4}-\d{2}", month):
            raise ValueError("月份必须采用 YYYY-MM 格式")
        try:
            datetime.strptime(month, "%Y-%m")
        except ValueError as exc:
            raise ValueError("月份必须采用有效的 YYYY-MM 格式") from exc
        rows = self.database.list_records(
            date_from=f"{month}-01", date_to=f"{month}-31"
        )
        path = Path(target)
        workbook = Workbook()
        detail = workbook.active
        detail.title = "工作量明细"
        self._write_detail_sheet(detail, rows)
        self._write_summary_sheet(workbook.create_sheet("人员汇总"), rows, "person_name_snapshot", "人员")
        self._write_summary_sheet(workbook.create_sheet("项目汇总"), rows, "project_name_snapshot", "项目")
        self._write_summary_sheet(workbook.create_sheet("小组汇总"), rows, "team_name_snapshot", "小组")
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path

    def export_ai(self, month: str, target_directory: str | Path) -> Path:
        if not re.fullmatch(r"\d{4}-\d{2}", month):
            raise ValueError("月份必须采用 YYYY-MM 格式")
        try:
            datetime.strptime(month, "%Y-%m")
        except ValueError as exc:
            raise ValueError("月份必须采用有效的 YYYY-MM 格式") from exc
        rows = self.database.list_records(
            date_from=f"{month}-01", date_to=f"{month}-31"
        )
        target = Path(target_directory) / f"AI_Workload_{month}"
        target.mkdir(parents=True, exist_ok=True)
        self._write_csv(target / "workload_records.csv", rows)
        self._write_csv(target / "persons.csv", self.database.list_persons())
        self._write_csv(target / "projects.csv", self.database.list_projects())
        self._write_csv(target / "quota_rules.csv", self.database.list_rules())
        summary = {
            "month": month,
            "record_count": len(rows),
            "workload_total": self._sum_score(rows),
            "by_person": self._summary_dict(rows, "person_name_snapshot"),
            "by_project": self._summary_dict(rows, "project_name_snapshot"),
            "by_team": self._summary_dict(rows, "team_name_snapshot"),
        }
        (target / "summary.json").write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return target

    @staticmethod
    def _detail_values(row: dict[str, Any]) -> list[Any]:
        return [
            row["work_date"], row["person_name_snapshot"], row["team_name_snapshot"],
            row["project_name_snapshot"], row["module_snapshot"], row["category_snapshot"],
            row["rule_name_snapshot"], row["rule_code_snapshot"], float(row["quantity"]),
            row["input_unit_snapshot"], float(row["standardized_quantity"]), row["unit_snapshot"],
            float(row["base_score_snapshot"]), float(row["workload_score"]), row["remark"],
        ]

    def _write_detail_sheet(self, sheet: Any, rows: list[dict[str, Any]]) -> None:
        sheet.append(DETAIL_HEADERS)
        for row in rows:
            sheet.append(self._detail_values(row))
        self._style_sheet(sheet, decimal_columns={9, 11, 13, 14})

    def _write_summary_sheet(
        self, sheet: Any, rows: list[dict[str, Any]], key: str, label: str
    ) -> None:
        grouped: dict[str, list[Decimal | int]] = defaultdict(lambda: [0, Decimal("0")])
        for row in rows:
            name = str(row.get(key) or "未分组")
            grouped[name][0] += 1
            grouped[name][1] += Decimal(str(row["workload_score"]))
        sheet.append([label, "记录数", "工作量"])
        for name in sorted(grouped, key=str.casefold):
            count, score = grouped[name]
            sheet.append([name, int(count), float(score)])
        sheet.append(["合计", len(rows), float(Decimal(self._sum_score(rows)))])
        self._style_sheet(sheet, decimal_columns={3})

    @staticmethod
    def _style_sheet(sheet: Any, decimal_columns: set[int]) -> None:
        dark = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.fill = dark
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(row, column).value or "") for row in range(1, sheet.max_row + 1)]
            width = min(max(max(map(len, values), default=8) + 2, 10), 34)
            sheet.column_dimensions[get_column_letter(column)].width = width
        for column in decimal_columns:
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column).number_format = "0.0000"

    @staticmethod
    def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
        if not rows:
            path.write_text("", encoding="utf-8-sig")
            return
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)

    @staticmethod
    def _sum_score(rows: list[dict[str, Any]]) -> str:
        total = sum((Decimal(str(row["workload_score"])) for row in rows), Decimal("0"))
        return format(total, "f")

    def _summary_dict(self, rows: list[dict[str, Any]], key: str) -> dict[str, str]:
        grouped: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
        for row in rows:
            grouped[str(row.get(key) or "未分组")] += Decimal(str(row["workload_score"]))
        return {name: format(value, "f") for name, value in sorted(grouped.items())}
