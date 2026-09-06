from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.database import Database
from app.services.calculator import ValidationError, calculate
from app.services.export_service import ExportService
from app.services.record_service import RecordService


class WorkloadMvpTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="workload_mvp_test_")
        self.root = Path(self.temp.name)
        self.database = Database(self.root / "data" / "workload.db")
        self.database.initialize()
        self.records = RecordService(self.database)
        self.exports = ExportService(self.database)
        self.person = next(row for row in self.database.list_persons(active_only=True) if row["name"] == "Employee02")
        self.project = next(row for row in self.database.list_projects(active_only=True) if row["code"] == "DEMO_PROJECT_05")

    def tearDown(self) -> None:
        self.temp.cleanup()

    def rule(self, code: str) -> dict:
        return next(row for row in self.database.list_rules() if row["rule_code"] == code)

    def add(self, code: str, quantity: str, **kwargs) -> int:
        return self.records.save(
            work_date=kwargs.pop("work_date", "2026-09-05"),
            person_id=self.person["id"], project_id=self.project["id"],
            rule_id=self.rule(code)["id"], quantity=quantity, **kwargs,
        )

    def test_initial_dictionaries_and_empty_history(self) -> None:
        self.assertEqual(len(self.database.list_teams(active_only=True)), 4)
        self.assertEqual(len(self.database.list_persons(active_only=True)), 15)
        self.assertEqual(len(self.database.list_projects(active_only=True)), 8)
        self.assertEqual(len(self.database.list_rules(active_only=True)), 25)
        self.assertEqual(self.database.count_records(), 0)

    def test_new_project_is_immediately_available(self) -> None:
        project_id = self.database.add_project("NEW", "NEW PROJECT")
        active_ids = {row["id"] for row in self.database.list_projects(active_only=True)}
        self.assertIn(project_id, active_ids)

    def test_standard_calculation_and_length_conversion(self) -> None:
        pump = calculate(self.rule("M1-03-D"), "6")
        rack = calculate(self.rule("M1-04-A"), "240")
        self.assertEqual(str(pump.workload_score), "3.600000")
        self.assertEqual(str(rack.standardized_quantity), "2.400000")
        self.assertEqual(str(rack.workload_score), "6.000000")

    def test_special_rules_require_human_basis(self) -> None:
        with self.assertRaisesRegex(ValidationError, "本次基准分"):
            calculate(self.rule("M1-03-E"), "1", "")
        with self.assertRaisesRegex(ValidationError, "人工确认依据"):
            self.add("M1-03-E", "1", manual_base_score="3.5", remark="")
        record_id = self.add("M1-03-E", "1", manual_base_score="3.5", remark="会议确认")
        self.assertEqual(self.database.get_record(record_id)["workload_score"], "3.5")
        linked_id = self.add("M2-04-A", "10", remark="关联范围经负责人确认")
        self.assertEqual(self.database.get_record(linked_id)["workload_score"], "2.5")

    def test_manual_standardized_rule_is_not_guessed(self) -> None:
        with self.assertRaisesRegex(ValidationError, "人工确认依据"):
            self.add("M1-01-A", "2.5", remark="")
        record_id = self.add("M1-01-A", "2.5", remark="负责人确认标准化数量")
        record = self.database.get_record(record_id)
        self.assertEqual(record["standardized_quantity"], "2.5")
        self.assertEqual(record["workload_score"], "1.875")

    def test_rule_change_does_not_change_historical_snapshot(self) -> None:
        first_id = self.add("M1-03-B", "3")
        rule = self.rule("M1-03-B")
        updated = dict(rule)
        updated["base_score"] = "1.00"
        self.database.update_rule(rule["id"], updated)
        old = self.database.get_record(first_id)
        self.assertEqual(old["base_score_snapshot"], "0.9")
        self.assertEqual(old["workload_score"], "2.7")
        second_id = self.add("M1-03-B", "3")
        new = self.database.get_record(second_id)
        self.assertEqual(new["base_score_snapshot"], "1")
        self.assertEqual(new["workload_score"], "3")

    def test_stopping_project_keeps_history_but_blocks_new_records(self) -> None:
        record_id = self.add("M1-03-D", "2")
        self.database.set_active("projects", self.project["id"], False)
        self.assertIsNotNone(self.database.get_record(record_id))
        with self.assertRaisesRegex(ValidationError, "启用的项目"):
            self.add("M1-03-D", "2")

    def test_edit_rebuilds_snapshot_and_void_excludes_default_query(self) -> None:
        record_id = self.add("M1-03-D", "2")
        self.records.save(
            record_id=record_id, work_date="2026-09-06", person_id=self.person["id"],
            project_id=self.project["id"], rule_id=self.rule("M1-03-B")["id"],
            quantity="3", remark="已确认修改",
        )
        changed = self.database.get_record(record_id)
        self.assertEqual(changed["rule_code_snapshot"], "M1-03-B")
        self.assertEqual(changed["workload_score"], "2.7")
        self.database.void_record(record_id, "重复录入")
        self.assertEqual(self.database.list_records(), [])
        self.assertEqual(self.database.list_records(include_void=True)[0]["state"], "VOID")

    def test_validation_rejects_bad_date_and_negative_quantity(self) -> None:
        with self.assertRaisesRegex(ValidationError, "YYYY-MM-DD"):
            self.records.save(
                work_date="2026/09/05", person_id=self.person["id"], project_id=self.project["id"],
                rule_id=self.rule("M1-03-D")["id"], quantity="1",
            )
        with self.assertRaisesRegex(ValidationError, "不能为负数"):
            self.add("M1-03-D", "-1")

    def test_persistence_after_reopen(self) -> None:
        record_id = self.add("M1-03-D", "5")
        reopened = Database(self.database.path)
        reopened.initialize()
        self.assertEqual(reopened.get_record(record_id)["workload_score"], "3")

    def test_day_month_and_ai_exports_match_database(self) -> None:
        self.add("M1-03-D", "5", work_date="2026-09-05")
        self.add("M1-04-A", "200", work_date="2026-09-06")
        day_path = self.exports.export_day("2026-09-05", self.root / "day.xlsx")
        month_path = self.exports.export_month("2026-09", self.root / "month.xlsx")
        day_book = load_workbook(day_path, data_only=True, read_only=True)
        self.assertEqual(day_book["工作量明细"].max_row, 2)
        day_book.close()
        month_book = load_workbook(month_path, data_only=True, read_only=True)
        self.assertEqual(month_book.sheetnames, ["工作量明细", "人员汇总", "项目汇总", "小组汇总"])
        self.assertEqual(month_book["工作量明细"].max_row, 3)
        person_sheet = month_book["人员汇总"]
        self.assertEqual(person_sheet.cell(person_sheet.max_row, 3).value, 8.0)
        month_book.close()
        ai_path = self.exports.export_ai("2026-09", self.root)
        expected = {"workload_records.csv", "persons.csv", "projects.csv", "quota_rules.csv", "summary.json"}
        self.assertEqual({path.name for path in ai_path.iterdir()}, expected)
        summary = json.loads((ai_path / "summary.json").read_text(encoding="utf-8"))
        self.assertEqual(summary["record_count"], 2)
        self.assertEqual(summary["workload_total"], "8")


if __name__ == "__main__":
    unittest.main()
